from flask import (
    Blueprint,
    render_template,
    request,
    jsonify,
    session,
    redirect,
    url_for,
    flash,
    send_file
)

import os
import openpyxl

from werkzeug.utils import secure_filename
from openpyxl.styles import Font
from datetime import datetime

from database import get_db


admin_bp = Blueprint(
    "admin",
    __name__,
    url_prefix="/admin"
)


UPLOAD_FOLDER = "/tmp/uploads"

ALLOWED_EXTENSIONS = {
    "xlsx",
    "xls"
}


# ============================================================
# FILE CHECK
# ============================================================

def allowed_file(filename):

    return (
        "." in filename
        and filename.rsplit(
            ".",
            1
        )[1].lower()
        in ALLOWED_EXTENSIONS
    )


# ============================================================
# CREATE EXAM
# ============================================================

@admin_bp.route(
    "/create-exam",
    methods=["GET", "POST"]
)
def create_exam():

    if session.get("role") != "admin":

        return redirect(
            url_for("admin_login")
        )

    if request.method == "POST":

        title = request.form.get(
            "title"
        )

        description = request.form.get(
            "description"
        )

        duration = request.form.get(
            "duration_minutes"
        )

        passing_score = request.form.get(
            "passing_score"
        )

        db = get_db()

        response = (
            db.table("exams")
            .insert({

                "title": title,

                "description":
                    description,

                "duration_minutes":
                    int(duration),

                "passing_score":
                    float(passing_score)
                    if passing_score
                    else None,

                "is_active": True

            })
            .execute()
        )

        if not response.data:

            flash(
                "Could not create exam.",
                "error"
            )

            return redirect(
                request.url
            )

        exam_id = response.data[0]["id"]

        flash(
            "Exam created successfully!",
            "success"
        )

        return redirect(
            url_for(
                "admin.upload_questions",
                exam_id=exam_id
            )
        )

    return render_template(
        "admin/create_exam.html"
    )


# ============================================================
# UPLOAD QUESTIONS
# ============================================================

@admin_bp.route(
    "/upload-questions/<int:exam_id>",
    methods=["GET", "POST"]
)
def upload_questions(exam_id):

    if session.get("role") != "admin":

        return redirect(
            url_for("admin_login")
        )

    db = get_db()

    # --------------------------------------------------------
    # Get exam
    # --------------------------------------------------------

    exam_response = (
        db.table("exams")
        .select("*")
        .eq("id", exam_id)
        .limit(1)
        .execute()
    )

    exam = (
        exam_response.data[0]
        if exam_response.data
        else None
    )

    if not exam:

        return (
            "Exam not found",
            404
        )

    # --------------------------------------------------------
    # POST
    # --------------------------------------------------------

    if request.method == "POST":

        if "file" not in request.files:

            flash(
                "No file uploaded",
                "error"
            )

            return redirect(
                request.url
            )

        file = request.files["file"]

        if file.filename == "":

            flash(
                "No file selected",
                "error"
            )

            return redirect(
                request.url
            )

        if not allowed_file(
            file.filename
        ):

            flash(
                "Only Excel files (.xlsx, .xls) are allowed!",
                "error"
            )

            return redirect(
                request.url
            )

        filename = secure_filename(
            file.filename
        )

        filepath = os.path.join(
            UPLOAD_FOLDER,
            filename
        )

        try:

            file.save(filepath)

            workbook = openpyxl.load_workbook(
                filepath
            )

            sheet = workbook.active

            questions = []

            for row in sheet.iter_rows(
                min_row=2,
                values_only=True
            ):

                if not row[0]:
                    continue

                question = {

                    "exam_id":
                        exam_id,

                    "question_text":
                        str(row[0]),

                    "option_a":
                        str(row[1])
                        if row[1] is not None
                        else "",

                    "option_b":
                        str(row[2])
                        if row[2] is not None
                        else "",

                    "option_c":
                        str(row[3])
                        if row[3] is not None
                        else "",

                    "option_d":
                        str(row[4])
                        if row[4] is not None
                        else "",

                    "correct_answer":
                        str(row[5]).upper()
                        if row[5]
                        else "A",

                    "marks": 1

                }

                questions.append(
                    question
                )

            if not questions:

                flash(
                    "No questions found in Excel file.",
                    "error"
                )

                return redirect(
                    request.url
                )

            # ------------------------------------------------
            # Insert questions
            # ------------------------------------------------

            db.table(
                "questions"
            ).insert(
                questions
            ).execute()

            flash(
                f"{len(questions)} questions successfully uploaded!",
                "success"
            )

            return redirect(
                url_for(
                    "admin.view_results",
                    exam_id=exam_id
                )
            )

        except Exception as e:

            flash(
                f"Error processing file: {str(e)}",
                "error"
            )

            return redirect(
                request.url
            )

        finally:

            if os.path.exists(filepath):

                try:
                    os.remove(filepath)

                except Exception:
                    pass

    return render_template(
        "admin/upload_questions.html",
        exam=exam
    )


# ============================================================
# VIEW RESULTS
# ============================================================

@admin_bp.route(
    "/view-results/<int:exam_id>"
)
def view_results(exam_id):

    if session.get("role") != "admin":

        return redirect(
            url_for("admin_login")
        )

    db = get_db()

    # --------------------------------------------------------
    # Exam
    # --------------------------------------------------------

    exam_response = (
        db.table("exams")
        .select("*")
        .eq("id", exam_id)
        .limit(1)
        .execute()
    )

    exam = (
        exam_response.data[0]
        if exam_response.data
        else None
    )

    # --------------------------------------------------------
    # Attempts
    # --------------------------------------------------------

    attempt_response = (
        db.table("student_attempts")
        .select("*")
        .eq("exam_id", exam_id)
        .in_(
            "status",
            ["completed", "terminated"]
        )
        .order(
            "submitted_at",
            desc=True
        )
        .execute()
    )

    attempts = (
        attempt_response.data
        or []
    )

    results = []

    for attempt in attempts:

        user_response = (
            db.table("users")
            .select(
                "username,full_name"
            )
            .eq(
                "id",
                attempt["student_id"]
            )
            .limit(1)
            .execute()
        )

        user = (
            user_response.data[0]
            if user_response.data
            else {}
        )

        result = dict(attempt)

        result["username"] = user.get(
            "username"
        )

        result["full_name"] = user.get(
            "full_name"
        )

        results.append(result)

    return render_template(
        "admin/view_results.html",
        exam=exam,
        results=results
    )


# ============================================================
# VIEW LOGS
# ============================================================

@admin_bp.route(
    "/view-logs/<int:attempt_id>"
)
def view_logs(attempt_id):

    if session.get("role") != "admin":

        return redirect(
            url_for("admin_login")
        )

    db = get_db()

    # --------------------------------------------------------
    # Attempt
    # --------------------------------------------------------

    attempt_response = (
        db.table("student_attempts")
        .select("*")
        .eq("id", attempt_id)
        .limit(1)
        .execute()
    )

    if not attempt_response.data:

        return (
            "Attempt not found",
            404
        )

    attempt = attempt_response.data[0]

    # --------------------------------------------------------
    # Student
    # --------------------------------------------------------

    user_response = (
        db.table("users")
        .select(
            "username,full_name"
        )
        .eq(
            "id",
            attempt["student_id"]
        )
        .limit(1)
        .execute()
    )

    user = (
        user_response.data[0]
        if user_response.data
        else {}
    )

    # --------------------------------------------------------
    # Exam
    # --------------------------------------------------------

    exam_response = (
        db.table("exams")
        .select("title")
        .eq(
            "id",
            attempt["exam_id"]
        )
        .limit(1)
        .execute()
    )

    exam = (
        exam_response.data[0]
        if exam_response.data
        else {}
    )

    attempt["username"] = user.get(
        "username"
    )

    attempt["full_name"] = user.get(
        "full_name"
    )

    attempt["title"] = exam.get(
        "title"
    )

    # --------------------------------------------------------
    # Logs
    # --------------------------------------------------------

    logs_response = (
        db.table("monitoring_logs")
        .select("*")
        .eq(
            "attempt_id",
            attempt_id
        )
        .order(
            "timestamp",
            desc=False
        )
        .execute()
    )

    logs = (
        logs_response.data
        or []
    )

    # --------------------------------------------------------
    # Report
    # --------------------------------------------------------

    report_response = (
        db.table("exam_reports")
        .select("*")
        .eq(
            "attempt_id",
            attempt_id
        )
        .order(
            "timestamp",
            desc=True
        )
        .limit(1)
        .execute()
    )

    report = (
        report_response.data[0]
        if report_response.data
        else None
    )

    # --------------------------------------------------------
    # Calculate time
    # --------------------------------------------------------

    start_time = None
    end_time = None

    total_seconds = 0
    total_minutes = 0

    warning_logs = []

    def parse_time(ts):

        if not ts:
            return None

        try:

            if isinstance(ts, datetime):

                return ts

            return datetime.fromisoformat(
                str(ts).replace(
                    "Z",
                    "+00:00"
                )
            )

        except Exception:

            try:

                return datetime.strptime(
                    str(ts),
                    "%Y-%m-%d %H:%M:%S"
                )

            except Exception:

                return None

    if logs:

        start_time = parse_time(
            logs[0].get("timestamp")
        )

        end_time = parse_time(
            logs[-1].get("timestamp")
        )

        if start_time and end_time:

            total_seconds = round(
                (
                    end_time
                    - start_time
                ).total_seconds()
            )

            total_minutes = round(
                total_seconds / 60,
                1
            )

        warning_logs = [
            log
            for log in logs
            if log.get(
                "warning_issued"
            ) == 1
        ]

    return render_template(
        "admin/view_logs.html",
        attempt=attempt,
        logs=logs,
        warning_logs=warning_logs,
        report=report,
        total_seconds=total_seconds,
        total_minutes=total_minutes,
        start_time=start_time,
        end_time=end_time
    )


# ============================================================
# DELETE EXAM
# ============================================================

@admin_bp.route(
    "/delete-exam/<int:exam_id>"
)
def delete_exam(exam_id):

    if session.get("role") != "admin":

        return redirect(
            url_for("admin_login")
        )

    db = get_db()

    # Get attempts first
    attempt_response = (
        db.table("student_attempts")
        .select("id")
        .eq("exam_id", exam_id)
        .execute()
    )

    attempts = (
        attempt_response.data
        or []
    )

    attempt_ids = [
        attempt["id"]
        for attempt in attempts
    ]

    # Delete reports/logs
    for attempt_id in attempt_ids:

        db.table(
            "exam_reports"
        ).delete().eq(
            "attempt_id",
            attempt_id
        ).execute()

        db.table(
            "monitoring_logs"
        ).delete().eq(
            "attempt_id",
            attempt_id
        ).execute()

    # Delete attempts
    (
        db.table("student_attempts")
        .delete()
        .eq("exam_id", exam_id)
        .execute()
    )

    # Delete questions
    (
        db.table("questions")
        .delete()
        .eq("exam_id", exam_id)
        .execute()
    )

    # Delete exam
    (
        db.table("exams")
        .delete()
        .eq("id", exam_id)
        .execute()
    )

    flash(
        "Exam deleted successfully.",
        "success"
    )

    return redirect(
        url_for("admin_dashboard")
    )


# ============================================================
# EXPORT MONITORING
# ============================================================

@admin_bp.route(
    "/export-monitoring/<int:attempt_id>"
)
def export_monitoring(attempt_id):

    if session.get("role") != "admin":

        return redirect(
            url_for("admin_login")
        )

    db = get_db()

    response = (
        db.table("monitoring_logs")
        .select("*")
        .eq(
            "attempt_id",
            attempt_id
        )
        .order(
            "id",
            desc=False
        )
        .execute()
    )

    logs = (
        response.data
        or []
    )

    workbook = openpyxl.Workbook()

    sheet = workbook.active

    sheet.title = "Monitoring Logs"

    headers = [

        "ID",
        "Attempt ID",
        "Event Type",
        "Face Detected",
        "Gaze Direction",
        "Head Pose",
        "Warning Issued",
        "Details",
        "Timestamp"

    ]

    sheet.append(headers)

    for cell in sheet[1]:

        cell.font = Font(
            bold=True
        )

    for log in logs:

        sheet.append([

            log.get("id"),

            log.get(
                "attempt_id"
            ),

            log.get(
                "event_type"
            ),

            log.get(
                "face_detected"
            ),

            log.get(
                "gaze_direction"
            ),

            log.get(
                "head_pose"
            ),

            log.get(
                "warning_issued"
            ),

            log.get(
                "details"
            ),

            log.get(
                "timestamp"
            )

        ])

    filename = (
        f"/tmp/monitoring_attempt_"
        f"{attempt_id}.xlsx"
    )

    workbook.save(filename)

    return send_file(
        filename,
        as_attachment=True,
        download_name=(
            f"monitoring_attempt_"
            f"{attempt_id}.xlsx"
        )
    )